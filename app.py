import io
import os
from datetime import datetime, timedelta, timezone

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# ----------------------------------------------------------------------
# CONFIGURACIÓN GENERAL
# ----------------------------------------------------------------------
st.set_page_config(page_title="Registro de Entregas", page_icon="🚚", layout="centered")

NOMBRE_ARCHIVO = "registro_entregas.xlsx"
NOMBRE_HOJA = "Entregas"

# Si la app corre en un PC con Windows, el Excel se guarda directo en C:\bit.
# En Streamlit Cloud (Linux) se guarda en la carpeta de la app.
CARPETA_LOCAL = r"C:\bit"
ARCHIVO_EXCEL = NOMBRE_ARCHIVO
if os.name == "nt":
    try:
        os.makedirs(CARPETA_LOCAL, exist_ok=True)
        ARCHIVO_EXCEL = os.path.join(CARPETA_LOCAL, NOMBRE_ARCHIVO)
    except Exception:
        ARCHIVO_EXCEL = NOMBRE_ARCHIVO

# El logo debe estar en la misma carpeta que app.py (en la raíz del repositorio)
RUTA_LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logo.png")

ENCABEZADOS = [
    "Fecha de salida",
    "Hora de entrega",
    "Cliente",
    "Quién recepciona",
    "Conductor",
    "Transporte",
    "Guía / Factura",
]

CLIENTES_BASE = [
    "Viñedos de Aguirre",
    "Viña Casadonoso",
    "Viña Aromo",
    "Viña Casa Solis",
    "Viña Concha y Toro",
    "Viña Luis Felipe Edwards",
    "Transportes Colchagua",
    "Viña Kinast",
    "Viña Santa Alicia",
    "Viña Nuevo Mundo",
    "Molinet Impresores",
    "Viña Casas Patronales",
    "Folias Doña Cornelia",
    "Viña Araucano",
    "Viña Casas del Toqui",
    "Viña Top Wine",
    "Viña San Pedro", 
    "Viña Concha y Toro",
    "Viña Santa Rita",
    "Viña Undurraga",
    "Viña Almaviva",
    "Viña Cousiño Macul",
    "Viña Tarapacá",
    "Viña De Martino",
    "Viña Aquitania",
    "Viña Pérez Cruz",
    "Viña Concha y Toro",
    "Viña Santa Rita",
    "Viña Undurraga",
    "Viña Almaviva",
    "Viña Cousiño Macul",
    "Viña Tarapacá",
    "Viña De Martino",
    "Viña Aquitania",
    "Viña Pérez Cruz",
    "Viña Haras de Pirque",
    "Viña Montes",
    "Viña Viu Manent",
    "Viña Clos Apalta",
    "Viña Santa Cruz",
    "Viña Casa Silva",
    "Viña Lapostolle",
    "Viña Neyen",
    "Viña Koyle",
    "Viña Los Vascos",
    "Viña Laura Hartwig"


]

RECEPCIONISTAS_BASE = [
    "Camila Villasana",
    "Olga Gonzales",
    "Encargado Colchagua",
    "Eliza De Aguirre",
    "Javier Diaz",
    "Robinson Navarro",
    "Juan Sandoval",
    "Anselmo Bravo",
    "Leonardo Barrios",
    "Angel Marqui",
    "Juan Reyes",
    "Alexander Beltrán",
    "Ana Sanchez",
    "Retiro Materiales"
   
]

EMPRESAS_TRANSPORTE_BASE = [
    "Transportes Colchagua",
    "Starken",
    "Chilexpress",
    "Correos de Chile",
    "Blue Express",
    "Turbus Cargo",
    "Pullman Cargo"
]

OPCION_NUEVO = "➕ Agregar nuevo"
PLACEHOLDER_CLIENTE = "Seleccione cliente..."
PLACEHOLDER_RECEPCION = "Seleccione quién recepciona..."
OPCION_OTRA_EMPRESA = "Otra Empresa"
VEHICULO_INTERNO = "Vehículo Interno"

# Colores del diseño corporativo
COLOR_AZUL_MARINO = "1F3864"
COLOR_GRIS_CLARO = "F2F2F2"
COLOR_BLANCO = "FFFFFF"
COLOR_BORDE = "BFBFBF"


# ----------------------------------------------------------------------
# UTILIDADES
# ----------------------------------------------------------------------
def ahora_chile():
    """Devuelve la fecha y hora actual de Chile (el servidor de Streamlit Cloud usa UTC)."""
    try:
        from zoneinfo import ZoneInfo

        return datetime.now(ZoneInfo("America/Santiago"))
    except Exception:
        return datetime.now(timezone(timedelta(hours=-3)))


def aplicar_diseno(ws):
    """Aplica el diseño corporativo: cabecera azul marino y filas zebra gris claro."""
    borde = Side(style="thin", color=COLOR_BORDE)
    borde_celda = Border(left=borde, right=borde, top=borde, bottom=borde)

    relleno_cabecera = PatternFill(
        start_color=COLOR_AZUL_MARINO, end_color=COLOR_AZUL_MARINO, fill_type="solid"
    )
    relleno_gris = PatternFill(
        start_color=COLOR_GRIS_CLARO, end_color=COLOR_GRIS_CLARO, fill_type="solid"
    )
    relleno_blanco = PatternFill(
        start_color=COLOR_BLANCO, end_color=COLOR_BLANCO, fill_type="solid"
    )

    fuente_cabecera = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    fuente_normal = Font(name="Calibri", size=11, color="000000")

    # Cabecera
    for celda in ws[1]:
        celda.fill = relleno_cabecera
        celda.font = fuente_cabecera
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde_celda
    ws.row_dimensions[1].height = 28

    # Filas de datos con efecto zebra
    for idx_fila in range(2, ws.max_row + 1):
        relleno = relleno_gris if idx_fila % 2 == 0 else relleno_blanco
        for celda in ws[idx_fila]:
            celda.fill = relleno
            celda.font = fuente_normal
            celda.border = borde_celda
            celda.alignment = Alignment(horizontal="center", vertical="center")
        # Formato de fecha en la columna A
        ws.cell(row=idx_fila, column=1).number_format = "DD/MM/YYYY"
        ws.row_dimensions[idx_fila].height = 20

    # Ancho de columnas según contenido
    for idx_col in range(1, ws.max_column + 1):
        letra = get_column_letter(idx_col)
        largo_max = 0
        for celda in ws[letra]:
            if celda.value is not None:
                valor = (
                    celda.value.strftime("%d/%m/%Y")
                    if hasattr(celda.value, "strftime") and idx_col == 1
                    else str(celda.value)
                )
                largo_max = max(largo_max, len(valor))
        ws.column_dimensions[letra].width = min(max(largo_max + 4, 16), 45)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def crear_libro_nuevo():
    wb = Workbook()
    ws = wb.active
    ws.title = NOMBRE_HOJA
    ws.append(ENCABEZADOS)
    return wb


def abrir_libro():
    """Abre el Excel existente o crea uno nuevo si aún no existe."""
    if os.path.exists(ARCHIVO_EXCEL):
        try:
            wb = load_workbook(ARCHIVO_EXCEL)
            if NOMBRE_HOJA in wb.sheetnames:
                return wb, wb[NOMBRE_HOJA]
            ws = wb.create_sheet(NOMBRE_HOJA)
            ws.append(ENCABEZADOS)
            return wb, ws
        except Exception:
            pass
    wb = crear_libro_nuevo()
    return wb, wb[NOMBRE_HOJA]


def guardar_filas(filas):
    """Agrega las filas al Excel (una por guía) y reaplica el diseño."""
    wb, ws = abrir_libro()
    for fila in filas:
        ws.append(fila)
    aplicar_diseno(ws)
    wb.save(ARCHIVO_EXCEL)


def valores_guardados(indice_columna):
    """Lee valores únicos ya guardados en el Excel (para recordar 'nuevos' agregados antes)."""
    valores = []
    if not os.path.exists(ARCHIVO_EXCEL):
        return valores
    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)
        if NOMBRE_HOJA not in wb.sheetnames:
            return valores
        ws = wb[NOMBRE_HOJA]
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if len(fila) > indice_columna and fila[indice_columna]:
                valor = str(fila[indice_columna]).strip()
                if valor and valor not in valores:
                    valores.append(valor)
        wb.close()
    except Exception:
        pass
    return valores


def construir_opciones(base, extras_sesion, indice_columna, excluir=()):
    """Une la lista base + lo guardado en el Excel + lo agregado en esta sesión."""
    opciones = list(base)
    for valor in valores_guardados(indice_columna) + list(extras_sesion):
        if valor not in opciones and valor not in excluir:
            opciones.append(valor)
    return opciones


def leer_filas():
    """Devuelve todas las filas de datos guardadas en el Excel (sin la cabecera)."""
    if not os.path.exists(ARCHIVO_EXCEL):
        return []
    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)
        ws = wb[NOMBRE_HOJA]
        filas = [
            f for f in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in f)
        ]
        wb.close()
        return filas
    except Exception:
        return []


def generar_excel_presentacion():
    """Genera en memoria un Excel con diseño de presentación: logo, título, tabla
    corporativa con zebra, hoja de resumen por cliente y configuración de impresión."""
    filas = leer_filas()
    ahora = ahora_chile()

    azul = PatternFill(start_color=COLOR_AZUL_MARINO, end_color=COLOR_AZUL_MARINO, fill_type="solid")
    gris = PatternFill(start_color=COLOR_GRIS_CLARO, end_color=COLOR_GRIS_CLARO, fill_type="solid")
    blanco = PatternFill(start_color=COLOR_BLANCO, end_color=COLOR_BLANCO, fill_type="solid")
    celeste = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    lado = Side(style="thin", color=COLOR_BORDE)
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)

    wb = Workbook()
    ws = wb.active
    ws.title = "Registro de Entregas"
    ws.sheet_view.showGridLines = False

    n_cols = len(ENCABEZADOS)
    ultima_col = get_column_letter(n_cols)
    fila_cab = 6
    primera_dato = 7

    anchos = [18, 16, 32, 28, 18, 22, 24]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    # ---- Banner superior: logo + título ----
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 24

    logo_ok = False
    if os.path.exists(RUTA_LOGO):
        try:
            img = XLImage(RUTA_LOGO)
            alto = 64
            ratio = alto / img.height
            img.height = alto
            img.width = int(img.width * ratio)
            if img.width > 230:
                ratio = 230 / img.width
                img.width = 230
                img.height = int(img.height * ratio)
            ws.add_image(img, "A1")
            logo_ok = True
        except Exception:
            logo_ok = False

    col_titulo = "C" if logo_ok else "A"
    ws.merge_cells(f"{col_titulo}1:{ultima_col}2")
    celda_titulo = ws[f"{col_titulo}1"]
    celda_titulo.value = "REGISTRO DE ENTREGAS"
    celda_titulo.font = Font(name="Calibri", size=22, bold=True, color=COLOR_AZUL_MARINO)
    celda_titulo.alignment = Alignment(
        horizontal="right" if logo_ok else "left", vertical="center", indent=1
    )

    ws.merge_cells(f"{col_titulo}3:{ultima_col}3")
    celda_sub = ws[f"{col_titulo}3"]
    celda_sub.value = (
        f"Generado el {ahora.strftime('%d/%m/%Y %H:%M')} hrs   ·   "
        f"{len(filas)} guías registradas"
    )
    celda_sub.font = Font(name="Calibri", size=11, italic=True, color="595959")
    celda_sub.alignment = Alignment(
        horizontal="right" if logo_ok else "left", vertical="center", indent=1
    )

    # Línea azul marino separadora
    ws.row_dimensions[4].height = 5
    for c in range(1, n_cols + 1):
        ws.cell(row=4, column=c).fill = azul
    ws.row_dimensions[5].height = 10

    # ---- Cabecera de la tabla ----
    ws.row_dimensions[fila_cab].height = 30
    for j, titulo in enumerate(ENCABEZADOS, start=1):
        c = ws.cell(row=fila_cab, column=j, value=titulo)
        c.fill = azul
        c.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borde

    # ---- Datos con zebra ----
    for i, fila in enumerate(filas):
        r = primera_dato + i
        ws.row_dimensions[r].height = 21
        relleno = gris if i % 2 == 1 else blanco
        for j in range(1, n_cols + 1):
            valor = fila[j - 1] if j - 1 < len(fila) else None
            c = ws.cell(row=r, column=j, value=valor)
            c.fill = relleno
            c.font = Font(name="Calibri", size=11, color="000000")
            c.border = borde
            c.alignment = Alignment(horizontal="center", vertical="center")
            if j == 1:
                c.number_format = "DD/MM/YYYY"

    ultima_fila = max(primera_dato + len(filas) - 1, fila_cab)
    if filas:
        ws.auto_filter.ref = f"A{fila_cab}:{ultima_col}{ultima_fila}"
    ws.freeze_panes = f"A{primera_dato}"

    # ---- Configuración de impresión ----
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{fila_cab}:{fila_cab}"
    ws.print_area = f"A1:{ultima_col}{ultima_fila}"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.7
    ws.oddFooter.center.text = "Página &P de &N"
    ws.oddFooter.center.size = 9

    # ---- Hoja de resumen por cliente ----
    rs = wb.create_sheet("Resumen")
    rs.sheet_view.showGridLines = False
    rs.column_dimensions["A"].width = 38
    rs.column_dimensions["B"].width = 22

    rs.merge_cells("A1:B1")
    rs.row_dimensions[1].height = 32
    rs["A1"].value = "RESUMEN POR CLIENTE"
    rs["A1"].font = Font(name="Calibri", size=20, bold=True, color=COLOR_AZUL_MARINO)
    rs["A1"].alignment = Alignment(horizontal="left", vertical="center")

    rs.merge_cells("A2:B2")
    rs["A2"].value = f"Al {ahora.strftime('%d/%m/%Y %H:%M')} hrs"
    rs["A2"].font = Font(name="Calibri", size=11, italic=True, color="595959")

    rs.row_dimensions[4].height = 28
    for j, titulo in enumerate(["Cliente", "Guías entregadas"], start=1):
        c = rs.cell(row=4, column=j, value=titulo)
        c.fill = azul
        c.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borde

    conteo = {}
    for f in filas:
        cliente = (f[2] if len(f) > 2 and f[2] else "Sin cliente")
        conteo[cliente] = conteo.get(cliente, 0) + 1
    ordenado = sorted(conteo.items(), key=lambda x: (-x[1], x[0]))

    r = 5
    for i, (cliente, cantidad) in enumerate(ordenado):
        relleno = gris if i % 2 == 1 else blanco
        rs.row_dimensions[r].height = 21
        for j, valor in enumerate([cliente, cantidad], start=1):
            c = rs.cell(row=r, column=j, value=valor)
            c.fill = relleno
            c.font = Font(name="Calibri", size=11)
            c.border = borde
            c.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center", indent=1 if j == 1 else 0)
        r += 1

    rs.row_dimensions[r].height = 24
    for j, valor in enumerate(["TOTAL", len(filas)], start=1):
        c = rs.cell(row=r, column=j, value=valor)
        c.fill = celeste
        c.font = Font(name="Calibri", size=12, bold=True, color=COLOR_AZUL_MARINO)
        c.border = borde
        c.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center", indent=1 if j == 1 else 0)

    rs.page_setup.orientation = "portrait"
    rs.page_setup.paperSize = rs.PAPERSIZE_LETTER
    rs.page_setup.fitToWidth = 1
    rs.page_setup.fitToHeight = 0
    rs.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    rs.print_options.horizontalCentered = True
    rs.oddFooter.center.text = "Página &P de &N"
    rs.oddFooter.center.size = 9

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def cargar_tabla():
    if not os.path.exists(ARCHIVO_EXCEL):
        return pd.DataFrame(columns=ENCABEZADOS)
    try:
        wb = load_workbook(ARCHIVO_EXCEL, read_only=True)
        ws = wb[NOMBRE_HOJA]
        filas = list(ws.iter_rows(min_row=2, values_only=True))
        wb.close()
    except Exception:
        return pd.DataFrame(columns=ENCABEZADOS)

    df = pd.DataFrame(filas, columns=ENCABEZADOS)
    if not df.empty:
        df["Fecha de salida"] = df["Fecha de salida"].apply(
            lambda x: x.strftime("%d/%m/%Y") if hasattr(x, "strftime") else x
        )
    return df


# ----------------------------------------------------------------------
# ESTADO DE SESIÓN
# ----------------------------------------------------------------------
if "clientes_extra" not in st.session_state:
    st.session_state["clientes_extra"] = []
if "recepcionistas_extra" not in st.session_state:
    st.session_state["recepcionistas_extra"] = []
if "empresas_extra" not in st.session_state:
    st.session_state["empresas_extra"] = []
if "mensaje" not in st.session_state:
    st.session_state["mensaje"] = None


# ----------------------------------------------------------------------
# LÓGICA DE GUARDADO (se ejecuta al presionar el botón)
# ----------------------------------------------------------------------
def guardar_registro():
    ss = st.session_state

    # Cliente
    cliente = ss.get("cliente_sel")
    if cliente == OPCION_NUEVO:
        cliente = (ss.get("cliente_nuevo") or "").strip()
    elif cliente == PLACEHOLDER_CLIENTE:
        cliente = ""

    # Quién recepciona
    recepciona = ss.get("recepciona_sel")
    if recepciona == OPCION_NUEVO:
        recepciona = (ss.get("recepciona_nuevo") or "").strip()
    elif recepciona == PLACEHOLDER_RECEPCION:
        recepciona = ""

    # Conductor y transporte
    conductor = ss.get("conductor")
    if conductor == "Elizabeth":
        transporte = VEHICULO_INTERNO
    else:
        transporte = ss.get("transporte_sel")
        if transporte == OPCION_OTRA_EMPRESA:
            transporte = (ss.get("transporte_nuevo") or "").strip()

    # Guías: una por línea
    texto_guias = ss.get("guias_texto") or ""
    guias = [linea.strip() for linea in texto_guias.splitlines() if linea.strip()]

    # Validaciones
    errores = []
    if not cliente:
        errores.append("Debes indicar el cliente.")
    if not recepciona:
        errores.append("Debes indicar quién recepciona.")
    if not transporte:
        errores.append("Debes indicar la empresa de transporte.")
    if not guias:
        errores.append("Debes ingresar al menos una guía o factura.")

    if errores:
        ss["mensaje"] = ("error", " ".join(errores))
        return

    fecha_salida = ss.get("fecha_salida")
    ahora = ahora_chile()

    # Una fila independiente por cada guía, con su hora de entrega automática
    filas = []
    for guia in guias:
        if ss.get("hora_manual_activa") and ss.get("hora_manual"):
            hora_entrega = ss["hora_manual"].strftime("%H:%M:%S")
        else:
            hora_entrega = ahora_chile().strftime("%H:%M:%S")
        filas.append(
            [fecha_salida, hora_entrega, cliente, recepciona, conductor, transporte, guia]
        )

    try:
        guardar_filas(filas)
    except Exception as e:
        ss["mensaje"] = ("error", f"No se pudo guardar el Excel: {e}")
        return

    # Recordar nuevos valores en la sesión
    if ss.get("cliente_sel") == OPCION_NUEVO and cliente not in ss["clientes_extra"]:
        ss["clientes_extra"].append(cliente)
    if ss.get("recepciona_sel") == OPCION_NUEVO and recepciona not in ss["recepcionistas_extra"]:
        ss["recepcionistas_extra"].append(recepciona)
    if (
        conductor == "Empresa Externa"
        and ss.get("transporte_sel") == OPCION_OTRA_EMPRESA
        and transporte not in ss["empresas_extra"]
    ):
        ss["empresas_extra"].append(transporte)

    # Limpiar campos para el siguiente registro
    ss["guias_texto"] = ""
    ss["cliente_nuevo"] = ""
    ss["recepciona_nuevo"] = ""
    ss["transporte_nuevo"] = ""
    ss["hora_manual_activa"] = False
    ss["cliente_sel"] = PLACEHOLDER_CLIENTE
    ss["recepciona_sel"] = PLACEHOLDER_RECEPCION

    plural = "guía" if len(filas) == 1 else "guías"
    ss["mensaje"] = (
        "ok",
        f"Se guardaron {len(filas)} {plural} a las {ahora.strftime('%H:%M:%S')} hrs.",
    )


# ----------------------------------------------------------------------
# INTERFAZ
# ----------------------------------------------------------------------
if os.path.exists(RUTA_LOGO):
    col_izq, col_centro, col_der = st.columns([1, 2, 1])
    with col_centro:
        st.image(RUTA_LOGO, use_container_width=True)

st.title("🚚 Registro de Entregas")
st.caption("Completa los datos y guarda. Cada guía se registra en una fila independiente del Excel.")

# Mensaje del último guardado
if st.session_state["mensaje"]:
    tipo, texto = st.session_state["mensaje"]
    if tipo == "ok":
        st.success(texto)
    else:
        st.error(texto)
    st.session_state["mensaje"] = None

# 1. Fecha de salida
st.date_input("Fecha de salida", value=ahora_chile().date(), format="DD/MM/YYYY", key="fecha_salida")

# 2. Cliente
opciones_clientes = (
    [PLACEHOLDER_CLIENTE]
    + construir_opciones(CLIENTES_BASE, st.session_state["clientes_extra"], indice_columna=2)
    + [OPCION_NUEVO]
)
st.selectbox("Cliente", opciones_clientes, key="cliente_sel")
if st.session_state.get("cliente_sel") == OPCION_NUEVO:
    st.text_input("Nombre del nuevo cliente", key="cliente_nuevo")

# 3. Quién recepciona
opciones_recep = (
    [PLACEHOLDER_RECEPCION]
    + construir_opciones(
        RECEPCIONISTAS_BASE, st.session_state["recepcionistas_extra"], indice_columna=3
    )
    + [OPCION_NUEVO]
)
st.selectbox("Quién recepciona", opciones_recep, key="recepciona_sel")
if st.session_state.get("recepciona_sel") == OPCION_NUEVO:
    st.text_input("Nombre de quien recepciona", key="recepciona_nuevo")

# 4. Conductor
st.radio("Conductor", ["Elizabeth", "Empresa Externa"], horizontal=True, key="conductor")

# 5. Transporte
if st.session_state.get("conductor") == "Elizabeth":
    st.info(f"Transporte asignado automáticamente: **{VEHICULO_INTERNO}**")
else:
    opciones_transporte = construir_opciones(
        EMPRESAS_TRANSPORTE_BASE,
        st.session_state["empresas_extra"],
        indice_columna=5,
        excluir=(VEHICULO_INTERNO,),
    ) + [OPCION_OTRA_EMPRESA]
    st.selectbox("Empresa de transporte", opciones_transporte, key="transporte_sel")
    if st.session_state.get("transporte_sel") == OPCION_OTRA_EMPRESA:
        st.text_input("Nombre de la otra empresa", key="transporte_nuevo")

# Hora de entrega: automática por defecto, o manual si se activa
st.checkbox("Ingresar hora de entrega manualmente", key="hora_manual_activa")
if st.session_state.get("hora_manual_activa"):
    st.time_input(
        "Hora de entrega",
        value=ahora_chile().time().replace(second=0, microsecond=0),
        step=60,
        key="hora_manual",
    )
else:
    st.caption("La hora de entrega se registrará automáticamente al guardar.")

# 6. Guías / Facturas
st.text_area(
    "Guías / Facturas",
    height=200,
    placeholder="Escribe una guía por línea y presiona ENTER para la siguiente.\nEj:\n12345\n12346\n12347",
    key="guias_texto",
)

# Contador de guías detectadas
cantidad_guias = len([l for l in (st.session_state.get("guias_texto") or "").splitlines() if l.strip()])
if cantidad_guias:
    st.caption(f"Guías detectadas: {cantidad_guias}")

# 7. Guardar
st.button("💾 Guardar registro", type="primary", on_click=guardar_registro, use_container_width=True)

# ----------------------------------------------------------------------
# HISTORIAL Y DESCARGA
# ----------------------------------------------------------------------
st.divider()
st.subheader("📋 Registros guardados")

tabla = cargar_tabla()
if tabla.empty:
    st.write("Aún no hay registros.")
else:
    st.dataframe(tabla.iloc[::-1], use_container_width=True, hide_index=True)

if os.path.exists(ARCHIVO_EXCEL):
    st.download_button(
        "⬇️ Descargar Excel para presentación",
        data=generar_excel_presentacion(),
        file_name=f"Registro_Entregas_{ahora_chile().strftime('%Y-%m-%d')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
    if os.name == "nt":
        st.caption(f"💾 El registro se guarda automáticamente en: {ARCHIVO_EXCEL}")
    else:
        st.caption(
            "📁 Para que la descarga caiga en C:\\bit, en tu navegador configura la carpeta de "
            "descargas en C:\\bit (o activa «Preguntar dónde guardar cada archivo»)."
        )
        st.caption(
            "⚠️ En Streamlit Cloud el archivo en el servidor se reinicia cuando la app se reinicia "
            "o se vuelve a desplegar. Descarga el Excel con frecuencia."
        )
